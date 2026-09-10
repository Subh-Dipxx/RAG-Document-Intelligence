from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from backend.domain.models import Document, Page, ParsedDocument


class XlsxParser:
    def parse(self, file_path: str) -> ParsedDocument:
        path = Path(file_path)

        if not path.is_file():
            raise FileNotFoundError(file_path)

        workbook = load_workbook(
            filename=file_path,
            read_only=True,
            data_only=True,
        )

        pages: list[Page] = []

        try:
            for page_number, worksheet in enumerate(
                workbook.worksheets,
                start=1,
            ):
                lines: list[str] = [
                    f"Worksheet: {worksheet.title}",
                    "",
                ]

                for row in worksheet.iter_rows(values_only=True):
                    values = [
                        "" if value is None else str(value)
                        for value in row
                    ]

                    if any(value.strip() for value in values):
                        lines.append(" | ".join(values))

                pages.append(
                    Page(
                        number=page_number,
                        text="\n".join(lines),
                        metadata={
                            "parser": "xlsx",
                            "worksheet": worksheet.title,
                        },
                    )
                )
        finally:
            workbook.close()

        document = Document(
            id=path.stem,
            file_name=path.name,
            mime_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            pages=pages,
            metadata={
                "parser": "xlsx",
                "page_count": len(pages),
                "worksheets": [page.metadata.get("worksheet") for page in pages],
            },
        )

        return ParsedDocument(
            document=document,
            pages=pages,
            metadata={
                "parser": "xlsx",
                "page_count": len(pages),
            },
        )
