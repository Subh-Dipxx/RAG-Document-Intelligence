from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from backend.domain.models import Document, Page, ParsedDocument


class PptxParser:
    def parse(self, file_path: str) -> ParsedDocument:
        path = Path(file_path)

        if not path.is_file():
            raise FileNotFoundError(file_path)

        presentation = Presentation(file_path)

        pages: list[Page] = []

        for page_number, slide in enumerate(
            presentation.slides,
            start=1,
        ):
            lines: list[str] = []

            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text = shape.text.strip()
                    if text:
                        lines.append(text)

            pages.append(
                Page(
                    number=page_number,
                    text="\n".join(lines),
                    metadata={
                        "parser": "pptx",
                        "slide_number": page_number,
                    },
                )
            )

        document = Document(
            id=path.stem,
            file_name=path.name,
            mime_type=(
                "application/vnd.openxmlformats-officedocument."
                "presentationml.presentation"
            ),
            pages=pages,
            metadata={
                "parser": "pptx",
                "page_count": len(pages),
            },
        )

        return ParsedDocument(
            document=document,
            pages=pages,
            metadata={
                "parser": "pptx",
                "page_count": len(pages),
            },
        )
