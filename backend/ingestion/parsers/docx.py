from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument

from backend.domain.models import Document, Page, ParsedDocument


class DocxParser:
    def parse(self, file_path: str) -> ParsedDocument:
        path = Path(file_path)

        if not path.is_file():
            raise FileNotFoundError(file_path)

        source = DocxDocument(file_path)

        lines: list[str] = []

        for paragraph in source.paragraphs:
            text = paragraph.text.strip()
            if text:
                lines.append(text)

        # Basic table extraction for first-pass DOCX support.
        for table_index, table in enumerate(source.tables, start=1):
            lines.append("")
            lines.append(f"[Table {table_index}]")

            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells]
                lines.append(" | ".join(values))

        text = "\n".join(lines)

        page = Page(
            number=1,
            text=text,
            metadata={
                "parser": "docx",
                "page_count": 1,
            },
        )

        document = Document(
            id=path.stem,
            file_name=path.name,
            mime_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            pages=[page],
            metadata={
                "parser": "docx",
                "page_count": 1,
            },
        )

        return ParsedDocument(
            document=document,
            pages=[page],
            metadata={
                "parser": "docx",
                "page_count": 1,
            },
        )
